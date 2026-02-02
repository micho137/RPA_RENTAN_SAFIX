from __future__ import annotations

import json
import os
import re
import time
import logging
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Dict, Optional, Tuple

import pyautogui
from openpyxl import load_workbook
from pywinauto import Application, Desktop

from src.rentan.config.config import settings
from src.rentan.ui.overlay_status import StatusOverlay


logger = logging.getLogger(__name__)

PLACA_REGEX = re.compile(r"\bPLACA\b\s*[:\-]?\s*([A-Z0-9]{5,8})\b", re.IGNORECASE)


def log_print(message: str, level: str = "info") -> None:
    """
    Imprime en consola y registra en el log unificado.
    level: info | warning | error
    """
    print(message)
    if level == "warning":
        logger.warning(message)
    elif level == "error":
        logger.error(message)
    else:
        logger.info(message)


def _strip_quotes(s: str) -> str:
    if not s:
        return s
    s = s.strip()
    if (s.startswith("'") and s.endswith("'")) or (s.startswith('"') and s.endswith('"')):
        return s[1:-1].strip()
    return s


def normalize_placa(s: str) -> str:
    return re.sub(r"\s+", "", str(s or "").strip().upper())


def load_invoices_by_id(path: Path) -> Dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"No existe el agregado invoices_by_id: {path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("El archivo all_invoices_by_id.json no es un dict.")
    return data


def load_plate_catalog_from_excel(excel_path: Path) -> Dict[str, Dict[str, str]]:
    if not excel_path.exists():
        raise FileNotFoundError(f"No existe el Excel de placas: {excel_path}")

    wb = load_workbook(filename=str(excel_path), data_only=True)
    ws = wb.active

    headers: Dict[str, int] = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        h = str(cell.value or "").strip().upper()
        if h:
            headers[h] = col_idx

    required = ["PLACA", "CENTRO DE COSTOS", "INTERFACE"]
    missing = [r for r in required if r not in headers]
    if missing:
        raise ValueError(
            f"El Excel no tiene columnas requeridas {missing}. "
            f"Encontradas: {list(headers.keys())}"
        )

    catalog: Dict[str, Dict[str, str]] = {}
    for row_idx in range(2, ws.max_row + 1):
        placa_val = ws.cell(row=row_idx, column=headers["PLACA"]).value
        if not placa_val:
            continue

        placa = normalize_placa(placa_val)
        if not placa:
            continue

        centro_costos = ws.cell(row=row_idx, column=headers["CENTRO DE COSTOS"]).value
        interface = ws.cell(row=row_idx, column=headers["INTERFACE"]).value

        catalog[placa] = {
            "centro_costos": str(centro_costos or "").strip(),
            "interface": str(interface or "").strip(),
        }

    return catalog


def extract_doc_plate_and_total(
    invoice: Dict[str, Any],
    fallback_plate: str,
    key_fallback_doc_id: Optional[str] = None,
) -> Tuple[str, str, int]:
    doc_id = ((invoice.get("document") or {}).get("document_id")) or ""
    doc_id = str(doc_id).strip()

    if not doc_id and key_fallback_doc_id:
        doc_id = str(key_fallback_doc_id).strip()

    if not doc_id:
        raise ValueError("Factura sin document_id (ni en JSON ni en la key del agregado).")

    doc_id = doc_id.replace("-", "").strip()

    placa = invoice.get("placa") or (invoice.get("vehiculo") or {}).get("placa") or ""
    placa = str(placa).strip()

    if not placa:
        for item in (invoice.get("detalle") or []):
            desc = str((item or {}).get("descripcion") or "")
            m = PLACA_REGEX.search(desc)
            if m:
                placa = m.group(1).strip()
                break

    if not placa:
        m2 = PLACA_REGEX.search(str(fallback_plate))
        placa = (m2.group(1).strip() if m2 else str(fallback_plate).strip())

    placa = normalize_placa(placa)

    total_val = ((invoice.get("totales") or {}).get("total"))
    if total_val is None:
        raise ValueError(f"Factura {doc_id}: no trae totales.total")

    try:
        total_int = int(float(str(total_val).strip()))
    except Exception:
        raise ValueError(f"Factura {doc_id}: totales.total inválido: {total_val!r}")

    return doc_id, placa, total_int


@dataclass(frozen=True)
class SafixConfig:
    jnlp_path: Path
    main_window_title_re: str

    tesoreria_icon: Path
    valores_icon: Path
    z_icon: Path
    engranes_icon: Path

    user: str
    password: str
    nit: str
    xot_code: str
    got_code: str
    campo_84: str
    campo_05: str
    obl_code: str
    obl2_code: str

    form_ready_wait: float
    login_wait: float
    pyauto_pause: float
    write_interval: float
    wait_default: float
    wait_long: float
    wait_popup: float

    invoices_by_id_path: Path
    fallback_placa: str

    @staticmethod
    def from_settings() -> "SafixConfig":
        title_re = _strip_quotes(getattr(settings, "safix_window_title", ""))

        def _p(name: str) -> Path:
            v = getattr(settings, name, None)
            if not v:
                raise ValueError(f"Falta settings.{name}")
            return Path(v)

        return SafixConfig(
            jnlp_path=_p("safix_shortcut"),
            main_window_title_re=title_re,
            tesoreria_icon=_p("safix_tesoreria_icon"),
            valores_icon=_p("safix_valores_icon"),
            z_icon=_p("safix_z_icon"),
            engranes_icon=_p("safix_engranes_icon"),
            user=str(getattr(settings, "safix_user", "") or "").strip(),
            password=str(getattr(settings, "safix_pass", "") or "").strip(),
            nit=str(getattr(settings, "safix_nit", "") or "").strip(),
            xot_code=str(getattr(settings, "safix_xot_code", "") or "").strip(),
            got_code=str(getattr(settings, "safix_got_code", "") or "").strip(),
            campo_84=str(getattr(settings, "safix_campo_84", "") or "").strip(),
            campo_05=str(getattr(settings, "safix_campo_05", "") or "").strip(),
            obl_code=str(getattr(settings, "safix_obl_code", "OBL_EXCLU") or "OBL_EXCLU").strip(),
            obl2_code=str(getattr(settings, "safix_obl2_code", "OBL_ANTCON") or "OBL_ANTCON").strip(),
            form_ready_wait=float(getattr(settings, "safix_form_ready_wait", 0) or 0),
            login_wait=float(getattr(settings, "safix_login_wait", 0) or 0),
            pyauto_pause=float(getattr(settings, "safix_pyauto_pause", 0.0) or 0.0),
            write_interval=float(getattr(settings, "safix_write_interval", 0.05) or 0.05),
            wait_default=float(getattr(settings, "safix_wait_default", 0.7) or 0.7),
            wait_long=float(getattr(settings, "safix_wait_long", 1.2) or 1.2),
            wait_popup=float(getattr(settings, "safix_wait_popup", 1.6) or 1.6),
            invoices_by_id_path=_p("invoices_by_id_path"),
            fallback_placa=str(getattr(settings, "safix_placa", "") or "").strip(),
        )


class SafixAutomator:
    def __init__(self, cfg: SafixConfig) -> None:
        self.cfg = cfg
        self._main_handle: Optional[int] = None

        pyautogui.FAILSAFE = True
        pyautogui.PAUSE = self.cfg.pyauto_pause

    def _wait_for_window_win32(
        self,
        title_re: str,
        timeout: float = 60.0,
        interval: float = 1.0,
    ):
        start = time.time()
        desktop = Desktop(backend="win32")

        while True:
            try:
                win = desktop.window(title_re=title_re)
                if win.exists():
                    return win
            except Exception:
                pass

            if time.time() - start > timeout:
                raise TimeoutError(f"No apareció ventana con título que matchee: {title_re}")
            time.sleep(interval)

    def launch_and_focus_main(self):
        if not self.cfg.jnlp_path.exists():
            raise FileNotFoundError(f"No se encontró el JNLP: {self.cfg.jnlp_path}")

        os.startfile(str(self.cfg.jnlp_path))

        main_win = self._wait_for_window_win32(self.cfg.main_window_title_re, timeout=90.0)
        self._main_handle = int(main_win.handle)

        app = Application(backend="win32").connect(handle=main_win.handle)
        main_window = app.window(handle=main_win.handle)
        main_window.set_focus()
        return app, main_window

    def refocus_main(self):
        desktop = Desktop(backend="win32")

        if self._main_handle is not None:
            try:
                win = desktop.window(handle=self._main_handle)
                if win.exists():
                    app = Application(backend="win32").connect(handle=win.handle)
                    window = app.window(handle=win.handle)
                    window.set_focus()
                    return app, window
            except Exception:
                pass

        win = self._wait_for_window_win32(self.cfg.main_window_title_re, timeout=20.0)
        self._main_handle = int(win.handle)
        app = Application(backend="win32").connect(handle=win.handle)
        window = app.window(handle=win.handle)
        window.set_focus()
        return app, window

    def wait(self, t: Optional[float] = None):
        time.sleep(self.cfg.wait_default if t is None else t)

    def write_text(self, text: str):
        pyautogui.write(str(text), interval=self.cfg.write_interval)

    def write_text_safe(self, text: str, slow_min_interval: float = 0.10):
        pyautogui.hotkey("ctrl", "a")
        self.wait(0.2)
        pyautogui.press("backspace")
        self.wait(0.2)
        pyautogui.write(str(text), interval=max(self.cfg.write_interval, slow_min_interval))
        self.wait(0.6)

    def press_enter(self, n: int = 1, wait_each: Optional[float] = None):
        for _ in range(max(1, n)):
            pyautogui.press("enter")
            if wait_each is not None:
                self.wait(wait_each)

    def press_tab(self, n: int = 1, wait_each: Optional[float] = None):
        for _ in range(max(1, n)):
            pyautogui.press("tab")
            if wait_each is not None:
                self.wait(wait_each)

    def do_login(self):
        self.wait(1.2)
        self.write_text_safe(self.cfg.user)

        pyautogui.press("tab")
        self.wait(0.8)

        self.write_text_safe(self.cfg.password)
        pyautogui.press("enter")

    def click_image(self, image_path: str | Path, timeout: float = 40.0, interval: float = 1.0) -> bool:
        icon_path = Path(image_path)
        if not icon_path.exists():
            raise FileNotFoundError(f"Image not found: {icon_path}")

        start = time.time()
        while time.time() - start < timeout:
            pos = pyautogui.locateCenterOnScreen(str(icon_path))
            if pos:
                pyautogui.click(pos.x, pos.y)
                self.wait(self.cfg.wait_long)
                return True
            time.sleep(interval)

        raise RuntimeError(f"No se pudo ubicar imagen en pantalla: {icon_path}")

    def click_tesoreria(self):
        return self.click_image(self.cfg.tesoreria_icon)

    def alt_p_o_g(self):
        pyautogui.keyDown("alt")
        time.sleep(0.1)
        pyautogui.press("p")
        time.sleep(0.25)
        pyautogui.press("o")
        time.sleep(0.25)
        pyautogui.press("g")
        time.sleep(0.25)
        pyautogui.keyUp("alt")
        time.sleep(1.2)

    def escribir_modal(self, text_str: str):
        pyautogui.hotkey("ctrl", "l")
        self.wait(self.cfg.wait_default)

        self.press_tab(4, self.cfg.wait_default)
        pyautogui.press("right")
        self.wait(self.cfg.wait_default)

        self.write_text(text_str)
        self.wait(self.cfg.wait_long)

        self.press_enter(2, self.cfg.wait_long)

    def procesar_factura(
        self,
        document_id: str,
        placa: str,
        interface: str,
        centro_costos: str,
        peaje_total: int,
        status_cb=None
    ):
        _ = centro_costos

        self.wait(self.cfg.wait_long)

        self.write_text_safe(self.cfg.xot_code)
        pyautogui.press("enter")
        self.wait(self.cfg.wait_long * 1.5)

        self.write_text_safe(self.cfg.nit)
        pyautogui.press("enter")
        self.wait(self.cfg.wait_popup * 1.5)

        self.press_enter(1, self.cfg.wait_popup)
        self.press_tab(4, self.cfg.wait_long)

        self.write_text_safe(document_id)
        pyautogui.press("enter")
        self.wait(self.cfg.wait_long)

        self.press_enter(3, self.cfg.wait_long)

        self.escribir_modal(self.cfg.got_code)

        self.press_tab(1, self.cfg.wait_default)

        interface_to_write = (interface or "").strip() or self.cfg.campo_84
        self.write_text_safe(interface_to_write)
        self.press_enter(2, self.cfg.wait_long)

        self.write_text_safe(self.cfg.campo_05)
        pyautogui.press("enter")
        self.wait(self.cfg.wait_long)
        pyautogui.press("tab")
        self.wait(self.cfg.wait_default)

        self.press_tab(8, self.cfg.wait_default)

        self.write_text_safe(placa)
        self.wait(self.cfg.wait_long * 1.3)

        self.click_image(self.cfg.valores_icon)

        self.wait(self.cfg.wait_popup)

        self.write_text_safe(self.cfg.obl_code)
        pyautogui.press("enter")
        self.wait(self.cfg.wait_long)

        self.write_text_safe(str(peaje_total))
        self.press_enter(2, self.cfg.wait_long)
        self.wait(self.cfg.wait_long)

        self.write_text_safe(self.cfg.obl2_code)
        pyautogui.press("enter")
        self.wait(self.cfg.wait_long)

        self.write_text_safe(str(peaje_total))
        self.press_enter(2, self.cfg.wait_long)
        self.wait(self.cfg.wait_long)

        self.click_image(self.cfg.z_icon)
        self.wait(self.cfg.wait_long)

        self.press_tab(4, self.cfg.wait_long)
        self.escribir_modal(self.cfg.got_code)
        self.wait(self.cfg.wait_long)

        self.click_image(self.cfg.engranes_icon)
        self.wait(self.cfg.wait_long)
        self.press_enter(1, self.cfg.wait_long)
        self.wait(self.cfg.wait_default)

        pyautogui.press("esc")
        self.wait(0.3)
        self.refocus_main()
        self.wait(0.4)

    def bootstrap(self):
        self.launch_and_focus_main()
        time.sleep(self.cfg.form_ready_wait)

        self.refocus_main()
        self.wait(1.0)

        self.do_login()
        time.sleep(self.cfg.login_wait)

        self.refocus_main()
        self.wait(1.0)

        self.click_tesoreria()
        self.wait(self.cfg.wait_long * 2)

        self.alt_p_o_g()
        self.wait(self.cfg.wait_long)


def run_safix_from_aggregated_json() -> None:
    cfg = SafixConfig.from_settings()

    invoices_by_id = load_invoices_by_id(cfg.invoices_by_id_path)
    if not invoices_by_id:
        log_print("[SAFIX] No hay facturas. No se ejecuta.")
        return

    items = list(invoices_by_id.items())
    total_invoices = len(items)

    overlay = StatusOverlay(width=340, height=165)
    overlay.start()
    overlay.update(etapa="AIVO: RENTAN", current=0, total=total_invoices, extra="Iniciando SAFIX…")

    automator = SafixAutomator(cfg)

    log_print(f"[SAFIX] leyendo agregado: {cfg.invoices_by_id_path.resolve()}")
    log_print(f"[SAFIX] total facturas: {total_invoices}")

    automator.bootstrap()
    automator.refocus_main()
    automator.wait(1.0)

    ok = 0
    fail = 0

    for idx, (key, invoice) in enumerate(items, start=1):
        doc_id = ""
        placa = ""
        try:
            doc_id, placa, total = extract_doc_plate_and_total(
                invoice,
                fallback_plate=cfg.fallback_placa,
                key_fallback_doc_id=key,
            )

            overlay.update(
                etapa="AIVO: Procesando factura",
                document_id=doc_id,
                placa=placa,
                current=idx,
                total=total_invoices,
                extra="",
            )

            automator.refocus_main()
            automator.wait(0.6)

            automator.procesar_factura(
                document_id=doc_id,
                placa=placa,
                interface="",
                centro_costos="",
                peaje_total=total,
            )

            ok += 1

        except Exception as e:
            fail += 1
            overlay.update(
                etapa="AIVO: ERROR",
                document_id=doc_id,
                placa=placa,
                current=idx,
                total=total_invoices,
                extra=str(e),
            )
            log_print(f"[SAFIX][ERROR] key='{key}': {e}", level="error")
            try:
                automator.refocus_main()
                automator.wait(1.0)
            except Exception:
                pass
            continue

    overlay.update(
        etapa="AIVO: Finalizado",
        current=total_invoices,
        total=total_invoices,
        extra=f"OK={ok} FAIL={fail}",
    )
    log_print(f"[SAFIX] Finalizado. OK={ok} FAIL={fail}")


def run_safix_with_excel(
    excel_path: Path,
    aggregated_json_path: Optional[Path] = None,
    tracker=None,
) -> None:
    from datetime import datetime

    overlay = StatusOverlay(width=340, height=165)
    overlay.start()
    overlay.update(etapa="AIVO: RENTAN", extra="Abriendo SAFIX y preparando sesión…")

    cfg = SafixConfig.from_settings()
    if aggregated_json_path is not None:
        cfg = replace(cfg, invoices_by_id_path=aggregated_json_path)

    plate_catalog = load_plate_catalog_from_excel(excel_path)
    log_print(f"[SAFIX] catálogo placas cargado: {len(plate_catalog)} desde {excel_path.resolve()}")

    invoices_by_id = load_invoices_by_id(cfg.invoices_by_id_path)
    if not invoices_by_id:
        overlay.update(etapa="AIVO: RENTAN", extra="No hay facturas. No se ejecuta.")
        log_print("[SAFIX] No hay facturas. No se ejecuta.")
        return

    items = list(invoices_by_id.items())
    total_invoices = len(items)

    overlay.update(
        etapa="AIVO: RENTAN",
        document_id="",
        placa="",
        current=0,
        total=total_invoices,
        extra=f"Facturas detectadas: {total_invoices}",
    )

    automator = SafixAutomator(cfg)

    log_print(f"[SAFIX] leyendo agregado: {cfg.invoices_by_id_path.resolve()}")
    log_print(f"[SAFIX] total facturas: {total_invoices}")

    automator.bootstrap()
    automator.refocus_main()
    automator.wait(1.0)

    ok = 0
    fail = 0
    missing_plate = 0

    for idx, (key, invoice) in enumerate(items, start=1):
        doc_id = ""
        placa = ""
        total = None
        interface = ""
        centro_costos = ""

        t_start = datetime.now().isoformat(timespec="seconds")

        try:
            doc_id, placa, total = extract_doc_plate_and_total(
                invoice,
                fallback_plate=cfg.fallback_placa,
                key_fallback_doc_id=key,
            )

            row = plate_catalog.get(placa)
            if not row:
                missing_plate += 1
                interface = ""
                centro_costos = ""
                extra = f"placa no está en catálogo | total={total}"
            else:
                interface = row.get("interface", "")
                centro_costos = row.get("centro_costos", "")
                extra = f"interface={interface or '-'} | cc={centro_costos or '-'} | total={total}"

            overlay.update(
                etapa="AIVO: Procesando factura",
                document_id=doc_id,
                placa=placa,
                current=idx,
                total=total_invoices,
                extra=extra,
            )

            automator.refocus_main()
            automator.wait(0.6)

            automator.procesar_factura(
                document_id=doc_id,
                placa=placa,
                interface=interface,
                centro_costos=centro_costos,
                peaje_total=int(total),
            )

            t_end = datetime.now().isoformat(timespec="seconds")
            if tracker is not None:
                tracker.add_processed(
                    timestamp_start=t_start,
                    timestamp_end=t_end,
                    document_id=doc_id,
                    placa=placa,
                    total=int(total) if total is not None else None,
                    key=str(key),
                    status="OK",
                    error="",
                )

            ok += 1

        except Exception as e:
            fail += 1
            t_end = datetime.now().isoformat(timespec="seconds")

            overlay.update(
                etapa="AIVO: ERROR",
                document_id=doc_id,
                placa=placa,
                current=idx,
                total=total_invoices,
                extra=f"{type(e).__name__}: {e}",
            )

            if tracker is not None:
                tracker.add_processed(
                    timestamp_start=t_start,
                    timestamp_end=t_end,
                    document_id=doc_id,
                    placa=placa,
                    total=int(total) if isinstance(total, (int, float)) else None,
                    key=str(key),
                    status="FAIL",
                    error=str(e),
                )

            log_print(f"[SAFIX][ERROR] key='{key}': {e}", level="error")
            try:
                automator.refocus_main()
                automator.wait(1.0)
            except Exception:
                pass
            continue

    overlay.update(
        etapa="AIVO: Finalizado",
        document_id="",
        placa="",
        current=total_invoices,
        total=total_invoices,
        extra=f"OK={ok} FAIL={fail} missing_plate={missing_plate}",
    )

    log_print(f"[SAFIX] Finalizado. OK={ok} FAIL={fail} missing_plate={missing_plate}")
