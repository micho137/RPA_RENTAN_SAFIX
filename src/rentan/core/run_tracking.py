from __future__ import annotations

from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict, Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


def _now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _autosize(ws):
    for col in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col)
        for cell in ws[col_letter]:
            v = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(v))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 70)


@dataclass
class ProcessRow:
    timestamp_start: str
    timestamp_end: str
    document_id: str
    placa: str
    total: Optional[int]
    key: str
    status: str  # OK / FAIL / PENDIENTE_DOBLE_CC
    error: str = ""


class RunTracker:
    """
    Guarda el log de procesadas en:
      <output_dir>/logs/procesadas.xlsx

    - add_processed(): agrega a memoria
    - flush_processed(): persiste inmediatamente (append) en Excel
    - save(): persiste todo (sobrescribe)
    """

    def __init__(self, output_dir: Path) -> None:
        self.output_dir = Path(output_dir)
        self.logs_dir = self.output_dir / "logs"
        self.logs_dir.mkdir(parents=True, exist_ok=True)

        self.processed: List[ProcessRow] = []

        self.procesadas_path = self.logs_dir / "procesadas.xlsx"
        self._ensure_procesadas_file()

    # -------------------------
    # API principal
    # -------------------------
    def add_processed(
        self,
        *,
        timestamp_start: str,
        timestamp_end: str,
        document_id: str,
        placa: str,
        total: Optional[int],
        key: str,
        status: str,
        error: str = "",
        flush: bool = True,
    ) -> None:
        row = ProcessRow(
            timestamp_start=timestamp_start,
            timestamp_end=timestamp_end,
            document_id=document_id,
            placa=placa,
            total=total,
            key=key,
            status=status,
            error=error,
        )
        self.processed.append(row)

        if flush:
            self.flush_processed([row])

    def flush_processed(self, rows: List[ProcessRow]) -> None:
        """
        Inserta filas incrementalmente en procesadas.xlsx (append),
        para que el archivo se vaya generando mientras corre SAFIX.
        """
        if not rows:
            return

        wb = load_workbook(self.procesadas_path)
        ws = wb["procesadas"]

        for r in rows:
            d = asdict(r)
            ws.append([
                d.get("timestamp_start", ""),
                d.get("timestamp_end", ""),
                d.get("document_id", ""),
                d.get("placa", ""),
                d.get("total", ""),
                d.get("key", ""),
                d.get("status", ""),
                d.get("error", ""),
            ])

        _autosize(ws)
        wb.save(self.procesadas_path)

    def save(self) -> Dict[str, Path]:
        """
        Guarda TODO lo que haya en memoria, sobrescribiendo el archivo.
        Útil si quieres regenerar completo al final.
        """
        self._write_procesadas_overwrite([asdict(x) for x in self.processed])
        return {"procesadas": self.procesadas_path}

    # Alias para evitar el error RunTracker has no attribute Save
    def Save(self) -> Dict[str, Path]:
        return self.save()

    # -------------------------
    # Internos
    # -------------------------
    def _ensure_procesadas_file(self) -> None:
        if self.procesadas_path.exists():
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "procesadas"
        ws.append(["timestamp_start", "timestamp_end", "document_id", "placa", "total", "key", "status", "error"])
        _autosize(ws)
        wb.save(self.procesadas_path)

    def _write_procesadas_overwrite(self, rows: List[Dict[str, Any]]) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "procesadas"

        ws.append(["timestamp_start", "timestamp_end", "document_id", "placa", "total", "key", "status", "error"])

        for r in rows:
            ws.append([
                r.get("timestamp_start", ""),
                r.get("timestamp_end", ""),
                r.get("document_id", ""),
                r.get("placa", ""),
                r.get("total", ""),
                r.get("key", ""),
                r.get("status", ""),
                r.get("error", ""),
            ])

        _autosize(ws)
        wb.save(self.procesadas_path)
