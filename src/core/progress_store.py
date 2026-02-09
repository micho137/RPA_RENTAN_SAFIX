from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional, Set

from openpyxl import Workbook, load_workbook


def _now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


@dataclass
class ProgressRow:
    timestamp: str
    document_id: str
    placa: str
    status: str
    error: str = ""
    source_key: str = ""


class ProgressStore:
    """
    Persistencia global en src/logs para reanudación entre ejecuciones.
    - state JSON: lógica de resume
    - descargadas/procesadas XLSX: visibilidad operativa
    """

    def __init__(self, logs_dir: Path) -> None:
        self.logs_dir = Path(logs_dir)
        self.logs_dir.mkdir(parents=True, exist_ok=True)
        self.state_path = self.logs_dir / "resume_state.json"
        self.downloaded_xlsx = self.logs_dir / "facturas_descargadas.xlsx"
        self.processed_xlsx = self.logs_dir / "facturas_procesadas_global.xlsx"
        self._state = self._load_state()

    def _load_state(self) -> Dict[str, Dict[str, str]]:
        if not self.state_path.exists():
            return {"invoices": {}}
        try:
            data = json.loads(self.state_path.read_text(encoding="utf-8"))
            if not isinstance(data, dict):
                return {"invoices": {}}
            invoices = data.get("invoices")
            if not isinstance(invoices, dict):
                return {"invoices": {}}
            return {"invoices": invoices}
        except Exception:
            return {"invoices": {}}

    def _save_state(self) -> None:
        self.state_path.write_text(
            json.dumps(self._state, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def get_processed_ok_ids(self) -> Set[str]:
        invoices = self._state.get("invoices", {})
        out: Set[str] = set()
        for doc_id, payload in invoices.items():
            if str((payload or {}).get("status", "")).upper() == "OK":
                out.add(str(doc_id).strip())
        return out

    def register_downloaded(self, *, document_id: str, placa: str, source_key: str = "") -> None:
        doc_id = str(document_id or "").strip()
        if not doc_id:
            return

        invoices = self._state["invoices"]
        prev = invoices.get(doc_id, {})
        invoices[doc_id] = {
            "document_id": doc_id,
            "placa": str(placa or "").strip(),
            "status": str(prev.get("status", "DOWNLOADED")).strip() or "DOWNLOADED",
            "error": str(prev.get("error", "")).strip(),
            "source_key": str(source_key or prev.get("source_key", "")).strip(),
            "updated_at": _now_iso(),
        }
        self._save_state()
        self._append_xlsx(
            self.downloaded_xlsx,
            "descargadas",
            ProgressRow(
                timestamp=_now_iso(),
                document_id=doc_id,
                placa=str(placa or "").strip(),
                status="DOWNLOADED",
                source_key=str(source_key or ""),
            ),
        )

    def mark_processed(self, *, document_id: str, placa: str, status: str, error: str = "", source_key: str = "") -> None:
        doc_id = str(document_id or "").strip()
        if not doc_id:
            return

        invoices = self._state["invoices"]
        prev = invoices.get(doc_id, {})
        invoices[doc_id] = {
            "document_id": doc_id,
            "placa": str(placa or prev.get("placa", "")).strip(),
            "status": str(status or "").strip().upper() or "UNKNOWN",
            "error": str(error or "").strip(),
            "source_key": str(source_key or prev.get("source_key", "")).strip(),
            "updated_at": _now_iso(),
        }
        self._save_state()
        self._append_xlsx(
            self.processed_xlsx,
            "procesadas",
            ProgressRow(
                timestamp=_now_iso(),
                document_id=doc_id,
                placa=str(placa or prev.get("placa", "")).strip(),
                status=str(status or "").strip().upper() or "UNKNOWN",
                error=str(error or "").strip(),
                source_key=str(source_key or prev.get("source_key", "")).strip(),
            ),
        )

    def _append_xlsx(self, path: Path, sheet_name: str, row: ProgressRow) -> None:
        headers = ["timestamp", "document_id", "placa", "status", "error", "source_key"]
        values = [row.timestamp, row.document_id, row.placa, row.status, row.error, row.source_key]

        if path.exists():
            wb = load_workbook(filename=str(path))
            ws = wb.active
            if ws.max_row == 0 or (ws.max_row == 1 and ws.cell(1, 1).value is None):
                ws.append(headers)
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws.append(headers)

        ws.append(values)
        wb.save(path)
