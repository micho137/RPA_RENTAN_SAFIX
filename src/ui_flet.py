# ui_flet.py
import threading
from datetime import datetime, date
from pathlib import Path
import unicodedata

import flet as ft
from openpyxl import load_workbook

from src.workflows.invoice_pipeline import run_pipeline

DATE_FMT = "%d/%m/%Y"

ALLOWED_EXCEL_BASENAME = "Centros de Costos Vehiculos"
EXPECTED_HEADERS = ["N° VEHICULO", "PLACA", "UBICACIÓN", "CENTRO DE COSTOS", "INTERFACE"]


def main(page: ft.Page):
    page.title = "Parámetros de consulta"
    page.padding = 15
    page.scroll = None

    BASE_W, BASE_H = 720, 600
    EXP_W, EXP_H = 860, 720

    try:
        page.window.resizable = False
        page.window_width = BASE_W
        page.window_height = BASE_H
    except Exception:
        pass

    if hasattr(page, "window") and page.window is not None:
        try:
            page.window.width = BASE_W
            page.window.height = BASE_H
            page.window.resizable = False
        except Exception:
            pass

    def center_window_best_effort():
        try:
            page.window.center()
        except Exception:
            pass

    center_window_best_effort()

    def fmt_date(d):
        return d.strftime(DATE_FMT)

    def expand_window():
        try:
            if hasattr(page, "window") and page.window is not None:
                if (page.window.width, page.window.height) != (EXP_W, EXP_H):
                    page.window.width = EXP_W
                    page.window.height = EXP_H
                    center_window_best_effort()
                    page.update()
        except Exception:
            pass

    def restore_window():
        try:
            if hasattr(page, "window") and page.window is not None:
                if (page.window.width, page.window.height) != (BASE_W, BASE_H):
                    page.window.width = BASE_W
                    page.window.height = BASE_H
                    center_window_best_effort()
                    page.update()
        except Exception:
            pass

    # ---------------- VALIDACIÓN EXCEL ----------------
    def normalize_header(v) -> str:
        s = "" if v is None else str(v)
        s = s.replace("\u00a0", " ").strip()
        s = " ".join(s.split())
        s = s.replace("Nº", "N°").replace("No.", "N°").replace("No", "N°")
        s = s.upper()
        s = "".join(
            ch for ch in unicodedata.normalize("NFKD", s)
            if not unicodedata.combining(ch)
        )
        return s

    EXPECTED_HEADERS_NORM = [normalize_header(h) for h in EXPECTED_HEADERS]

    def validate_excel_file(path_str: str) -> tuple[bool, str]:
        p = Path(path_str)

        if not path_str.strip():
            return False, "Seleccione un archivo Excel."

        if not p.exists():
            return False, "El archivo seleccionado no existe."

        if p.suffix.lower() != ".xlsx":
            return False, "El archivo debe ser .xlsx para validar cabeceras (guárdalo como .xlsx)."

        if p.stem != ALLOWED_EXCEL_BASENAME:
            return False, "El archivo debe llamarse exactamente: 'Centros de Costos Vehiculos.xlsx'."

        try:
            wb = load_workbook(filename=str(p), read_only=True, data_only=True)
            ws = wb.active

            first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not first_row:
                return False, "El archivo no tiene cabeceras en la fila 1."

            headers = list(first_row[: len(EXPECTED_HEADERS)])
            headers_norm = [normalize_header(h) for h in headers]

            if headers_norm != EXPECTED_HEADERS_NORM:
                esperado = " | ".join(EXPECTED_HEADERS)
                encontrado = " | ".join("" if h is None else str(h) for h in headers)
                return (
                    False,
                    "Las cabeceras no coinciden.\n"
                    f"Esperado (A1:E1): {esperado}\n"
                    f"Encontrado (A1:E1): {encontrado}"
                )

            return True, ""

        except Exception as ex:
            return False, f"No se pudo leer el Excel para validar cabeceras. Detalle: {ex}"

    # ---------------- FILE PICKER ----------------
    selected_file_txt = ft.Text(value="", selectable=True)

    def on_file_picked(e):
        if e.files:
            selected_file_txt.value = e.files[0].path
        else:
            selected_file_txt.value = ""
        refresh()
        page.update()

    file_picker = ft.FilePicker(on_result=on_file_picked)
    page.overlay.append(file_picker)

    def open_file_picker(_):
        file_picker.pick_files(
            allow_multiple=False,
            allowed_extensions=["xlsx"],
        )

    from datetime import datetime, date

    def to_date(v) -> date:
        if isinstance(v, datetime):
            return v.date()
        return v  # ya es date

    # ---------------- FECHAS ----------------
    today = datetime.now().date()
    selected_start: date = today
    selected_end: date = today

    def clamp_range():
        nonlocal selected_start, selected_end
        if selected_end < selected_start:
            selected_end = selected_start

    def on_dp_desde_change(e):
        nonlocal selected_start
        if dp_desde.value:
            selected_start = to_date(dp_desde.value)
            clamp_range()
            refresh()
        restore_window()

    def on_dp_hasta_change(e):
        nonlocal selected_end
        if dp_hasta.value:
            selected_end = to_date(dp_hasta.value)
            clamp_range()
            refresh()
        restore_window()

    def build_datepicker(on_change_cb, on_dismiss_cb):
        kwargs = dict(
            first_date=date(2000, 1, 1),
            last_date=date(2100, 12, 31),
            on_change=on_change_cb,
            on_dismiss=on_dismiss_cb,
        )
        try:
            return ft.DatePicker(locale="es", **kwargs)
        except TypeError:
            try:
                return ft.DatePicker(locale="es-ES", **kwargs)
            except TypeError:
                return ft.DatePicker(**kwargs)

    dp_desde = build_datepicker(on_dp_desde_change, lambda e: restore_window())
    dp_hasta = build_datepicker(on_dp_hasta_change, lambda e: restore_window())

    page.overlay.append(dp_desde)
    page.overlay.append(dp_hasta)

    def open_desde(_):
        expand_window()
        page.open(dp_desde)

    def open_hasta(_):
        expand_window()
        page.open(dp_hasta)

    # ---------------- CONTROLES ----------------
    only_unread_cb = ft.Checkbox(label="Solo no leídos", value=False)
    mark_as_read_cb = ft.Checkbox(label="Marcar como leídos", value=False)
    move_to_processed_cb = ft.Checkbox(label="Mover a procesados", value=False)

    status = ft.Text(value="", color=ft.Colors.RED_700)

    # ---------------- BOTONES ----------------
    load_excel_btn = ft.ElevatedButton(
        "Cargar archivo Excel",
        icon=ft.Icons.UPLOAD_FILE,
        on_click=open_file_picker,
    )

    run_btn = ft.ElevatedButton(
        "Ejecutar AIVO",
        icon=ft.Icons.PLAY_ARROW,
    )
    run_btn.disabled = True

    def set_status(msg, is_error=False):
        status.value = msg
        status.color = ft.Colors.RED_700 if is_error else ft.Colors.GREEN_700
        page.update()

    # ---------------- REFRESH UI ----------------
    def refresh(_=None):
        excel_path = (selected_file_txt.value or "").strip()
        if not excel_path:
            run_btn.disabled = True
            status.value = ""
            page.update()
            return

        ok, err = validate_excel_file(excel_path)
        if not ok:
            run_btn.disabled = True
            set_status(err, is_error=True)
        else:
            run_btn.disabled = False
            status.value = ""
            page.update()

    only_unread_cb.on_change = refresh
    mark_as_read_cb.on_change = refresh
    move_to_processed_cb.on_change = refresh

    # ---------------- EJECUCIÓN ----------------
    def run_job(excel_path_str: str):
        try:
            set_status("Ejecutando pipeline (descarga / extracción / agregación)...")

            # ✅ flags calculados desde UI
            days_back = max(0, (selected_end - selected_start).days)
            only_unread = bool(only_unread_cb.value)
            mark_as_read = bool(mark_as_read_cb.value)
            move_to_processed = bool(move_to_processed_cb.value)

            run_pipeline(
                output_dir=Path("./output"),
                lang="spa",
                dpi=300,
                aggregate_by_id=True,
                excel_path=Path(excel_path_str),
                run_safix=True,

                # ✅ NUEVO: flags dinámicos
                only_unread=only_unread,
                days_back=days_back,
                mark_as_read=mark_as_read,
                move_to_processed=move_to_processed,
            )

            set_status("Proceso finalizado correctamente. Pipeline + SAFIX ejecutados.")

        except Exception as ex:
            set_status(f"Error: {ex}", is_error=True)
        finally:
            run_btn.disabled = False
            page.update()

    def on_run_click(_):
        excel_path_str = (selected_file_txt.value or "").strip()

        ok, err = validate_excel_file(excel_path_str)
        if not ok:
            set_status(err, is_error=True)
            return

        run_btn.disabled = True
        page.update()

        threading.Thread(
            target=run_job,
            args=(excel_path_str,),
            daemon=True,
        ).start()

    run_btn.on_click = on_run_click

    # ---------------- LAYOUT ----------------
    page.add(
        ft.Column(
            spacing=14,
            controls=[
                ft.Text("Configuración de consulta", size=20, weight=ft.FontWeight.BOLD),
                ft.Container(
                    content=ft.Column(
                        spacing=10,
                        controls=[
                            ft.Text("Fechas de consulta", weight=ft.FontWeight.W_600),
                            ft.Row(
                                controls=[
                                    ft.TextField(
                                        label="Fecha inicial",
                                        value=fmt_date(selected_start),
                                        read_only=True,
                                        width=280,
                                        dense=True,
                                    ),
                                    ft.IconButton(
                                        icon=ft.Icons.CALENDAR_MONTH,
                                        tooltip="Elegir fecha inicial",
                                        on_click=open_desde,
                                    ),
                                ]
                            ),
                            ft.Row(
                                controls=[
                                    ft.TextField(
                                        label="Fecha final",
                                        value=fmt_date(selected_end),
                                        read_only=True,
                                        width=280,
                                        dense=True,
                                    ),
                                    ft.IconButton(
                                        icon=ft.Icons.CALENDAR_MONTH,
                                        tooltip="Elegir fecha final",
                                        on_click=open_hasta,
                                    ),
                                ]
                            ),
                            ft.Divider(),
                            only_unread_cb,
                            mark_as_read_cb,
                            move_to_processed_cb,
                        ],
                    ),
                    padding=16,
                    border_radius=10,
                    border=ft.border.all(1, ft.Colors.GREY_300),
                ),
                ft.Divider(),
                load_excel_btn,
                selected_file_txt,
                run_btn,
                status,
            ],
        )
    )

    refresh()


if __name__ == "__main__":
    ft.app(target=main, view=ft.AppView.FLET_APP)
