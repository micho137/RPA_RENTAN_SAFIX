import threading
from datetime import date, datetime
from pathlib import Path
import unicodedata

import flet as ft
from openpyxl import load_workbook

from src.workflows.invoice_pipeline import run_pipeline

ALLOWED_EXCEL_BASENAME = "Centros de Costos Vehiculos"
EXPECTED_HEADERS = ["N VEHICULO", "PLACA", "UBICACION", "CENTRO DE COSTOS", "INTERFACE", "DOBLE CC"]


def main(page: ft.Page):
    from src.config import settings
    from src.core.logging_config import init_logging, redirect_std_streams

    logger = init_logging(settings.log_dir)
    redirect_std_streams(logger)

    page.title = "AIVO - Automatizacion SAFIX"
    page.padding = 16
    page.scroll = ft.ScrollMode.AUTO

    if hasattr(page, "window") and page.window is not None:
        try:
            page.window.width = 620
            page.window.height = 420
            page.window.min_width = 540
            page.window.min_height = 360
            page.window.resizable = True
            page.window.center()
        except Exception:
            pass

    # Parametros fijos solicitados
    only_unread_default = True
    mark_as_read_default = True
    days_back_default = 0

    def normalize_header(v) -> str:
        s = "" if v is None else str(v)
        s = s.replace("\xa0", " ").strip()
        s = " ".join(s.split())
        s = s.upper()
        s = "".join(ch for ch in unicodedata.normalize("NFKD", s) if not unicodedata.combining(ch))
        s = s.replace("N?", "N").replace("N?", "N").replace("NO.", "N").replace("NO", "N")
        s = s.replace("?", "")
        return s

    expected_headers_norm = [normalize_header(h) for h in EXPECTED_HEADERS]

    def validate_excel_file(path_str: str) -> tuple[bool, str]:
        pth = Path(path_str)

        if not path_str.strip():
            return False, "Seleccione un archivo Excel."
        if not pth.exists():
            return False, "El archivo seleccionado no existe."
        if pth.suffix.lower() != ".xlsx":
            return False, "El archivo debe ser .xlsx para validar cabeceras."
        if pth.stem != ALLOWED_EXCEL_BASENAME:
            return False, "El archivo debe llamarse exactamente: Centros de Costos Vehiculos.xlsx"

        try:
            wb = load_workbook(filename=str(pth), read_only=True, data_only=True)
            ws = wb.active
            first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not first_row:
                return False, "El archivo no tiene cabeceras en la fila 1."

            headers = list(first_row[: len(EXPECTED_HEADERS)])
            headers_norm = [normalize_header(h) for h in headers]
            if headers_norm != expected_headers_norm:
                esperado = " | ".join(EXPECTED_HEADERS)
                encontrado = " | ".join("" if h is None else str(h) for h in headers)
                return False, f"Cabeceras invalidas. Esperado: {esperado} / Encontrado: {encontrado}"

            return True, ""
        except Exception as ex:
            return False, f"No se pudo leer el Excel. Detalle: {ex}"

    selected_file_txt = ft.Text(value="", selectable=True, text_align=ft.TextAlign.CENTER)
    date_from_tf = ft.TextField(
        label="Desde (YYYY-MM-DD)",
        hint_text="2026-02-20",
        width=260,
    )
    date_to_tf = ft.TextField(
        label="Hasta (YYYY-MM-DD)",
        hint_text="2026-02-25",
        width=260,
    )
    status = ft.TextField(
        value="",
        read_only=True,
        multiline=True,
        min_lines=3,
        max_lines=10,
        width=560,
        text_size=13,
    )

    run_btn = ft.ElevatedButton("Ejecutar AIVO", icon=ft.Icons.PLAY_ARROW, disabled=True)

    def set_status(msg: str, is_error: bool = False):
        status.value = msg
        status.color = ft.Colors.RED_700 if is_error else ft.Colors.GREEN_700
        page.update()

    def parse_date_or_none(raw: str) -> date | None:
        value = (raw or "").strip()
        if not value:
            return None
        return datetime.strptime(value, "%Y-%m-%d").date()

    def validate_date_inputs() -> tuple[bool, str]:
        raw_from = (date_from_tf.value or "").strip()
        raw_to = (date_to_tf.value or "").strip()

        if bool(raw_from) ^ bool(raw_to):
            return False, "Debe diligenciar ambas fechas: Desde y Hasta."

        if not raw_from and not raw_to:
            return True, ""

        try:
            d_from = parse_date_or_none(raw_from)
            d_to = parse_date_or_none(raw_to)
        except ValueError:
            return False, "Formato de fecha invalido. Use YYYY-MM-DD."

        if d_from is None or d_to is None:
            return False, "Debe diligenciar ambas fechas: Desde y Hasta."
        if d_to < d_from:
            return False, "La fecha Hasta no puede ser menor que Desde."
        return True, ""

    def refresh(_=None):
        excel_path = (selected_file_txt.value or "").strip()
        if not excel_path:
            run_btn.disabled = True
            status.value = ""
            page.update()
            return

        ok, err = validate_excel_file(excel_path)
        if not ok:
            run_btn.disabled = True
            set_status(err, is_error=True)
            return

        ok_dates, err_dates = validate_date_inputs()
        if not ok_dates:
            run_btn.disabled = True
            set_status(err_dates, is_error=True)
        else:
            run_btn.disabled = False
            status.value = ""
            page.update()

    def run_job(excel_path_str: str):
        try:
            date_from = parse_date_or_none((date_from_tf.value or "").strip())
            date_to = parse_date_or_none((date_to_tf.value or "").strip())

            set_status("Ejecutando pipeline (descarga, extraccion, agregacion y SAFIX)...")
            run_pipeline(
                output_dir=None,
                lang="spa",
                dpi=300,
                aggregate_by_id=False,
                excel_path=Path(excel_path_str),
                run_safix=True,
                only_unread=only_unread_default,
                days_back=days_back_default,
                date_from=date_from,
                date_to=date_to,
                mark_as_read=mark_as_read_default,
            )
            set_status("Proceso finalizado correctamente. Pipeline + SAFIX ejecutados.")
        except Exception as ex:
            set_status(f"Error: {ex}", is_error=True)
        finally:
            run_btn.disabled = False
            page.update()
            try:
                if hasattr(page, "window") and page.window is not None:
                    page.window.minimized = False
                    page.window.focus()
                    page.update()
            except Exception:
                pass

    def on_run_click(_):
        excel_path_str = (selected_file_txt.value or "").strip()
        ok, err = validate_excel_file(excel_path_str)
        if not ok:
            set_status(err, is_error=True)
            return
        ok_dates, err_dates = validate_date_inputs()
        if not ok_dates:
            set_status(err_dates, is_error=True)
            return

        run_btn.disabled = True
        page.update()

        threading.Thread(target=run_job, args=(excel_path_str,), daemon=True).start()

    run_btn.on_click = on_run_click

    def on_file_picked(e: ft.FilePickerResultEvent):
        if e.files:
            selected_file_txt.value = e.files[0].path
        else:
            selected_file_txt.value = ""
        refresh()
        page.update()

    file_picker = ft.FilePicker(on_result=on_file_picked)
    page.overlay.append(file_picker)

    load_excel_btn = ft.ElevatedButton(
        "Cargar archivo Excel",
        icon=ft.Icons.UPLOAD_FILE,
        on_click=lambda _: file_picker.pick_files(allow_multiple=False, allowed_extensions=["xlsx"]),
    )

    page.add(
        ft.Container(
            alignment=ft.alignment.center,
            expand=True,
            content=ft.Column(
                alignment=ft.MainAxisAlignment.CENTER,
                horizontal_alignment=ft.CrossAxisAlignment.CENTER,
                spacing=14,
                controls=[
                    ft.Text("Carga de archivo", size=22, weight=ft.FontWeight.BOLD, text_align=ft.TextAlign.CENTER),
                    load_excel_btn,
                    selected_file_txt,
                    ft.Row(
                        alignment=ft.MainAxisAlignment.CENTER,
                        controls=[date_from_tf, date_to_tf],
                    ),
                    run_btn,
                    status,
                ],
            ),
        )
    )

    date_from_tf.on_change = refresh
    date_to_tf.on_change = refresh
    refresh()


if __name__ == "__main__":
    ft.app(target=main, view=ft.AppView.FLET_APP)
